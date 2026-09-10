"""Tests: sections that repeat per drug substance (3.2.S), P13.

Until P13 a registered section produced exactly one document, so a section
number was a sufficient identity. 3.2.S breaks that -- it is repeated per
drug substance -- and these tests pin the three things that repetition has
to get right: identity, placement, and content.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.storage import InMemoryStorageClient
from app.ctd.structure import folder_for_section_instance
from app.models import Base, SpecificationTest
import app.validation.rules  # noqa: F401  registers rules
from app.seed.ampiclox import build_ampiclox
from app.seed.examox import build_examox
from app.templating.instances import drug_substance_info, expand_sections, slugify_subject
from app.templating.render import render_section
from app.validation.engine import run_all


def _load(builder):
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)
    project = builder(buggy=False)
    session.add(project)
    session.commit()
    session.refresh(project)
    return project


# ---- identity ---------------------------------------------------------------


def test_a_combination_product_owes_one_copy_of_each_3_2_s_section_per_active():
    keys = [i.key for i in expand_sections(_load(build_ampiclox))]
    assert "3.2.S.1-ampicillin" in keys
    assert "3.2.S.1-cloxacillin" in keys
    assert "3.2.S.4.1-ampicillin" in keys
    assert "3.2.S.4.1-cloxacillin" in keys
    # and the bare number is never itself a document
    assert "3.2.S.1" not in keys


def test_non_repeating_sections_keep_exactly_the_identity_they_already_had():
    """The compatibility guarantee that let P13 skip a data migration: every
    pre-P13 section's key is still its bare number, so stored narratives and
    persisted sequence leaves keep resolving.

    P24d renumbered ONE of them, 1.2 -> 1.2.2, and deliberately: 1.2 is a
    heading in the target and a leaf was sitting at it (see the registry's
    note). That is the exception this test now records rather than hides --
    the guarantee is "identity is stable unless the contract says it was
    wrong", which is a different and more honest promise than "identity
    never changes".
    """
    keys = [i.key for i in expand_sections(_load(build_examox))]
    assert {"1.0", "1.2.2", "2.3", "3.2.P.1", "3.2.P.8.1"} <= set(keys)
    assert "1.2" not in keys


def test_titles_name_the_substance():
    """Two PDFs both titled "Specification" is a bookmark list an assessor
    cannot navigate."""
    titles = {i.key: i.title for i in expand_sections(_load(build_ampiclox))}
    assert titles["3.2.S.4.1-cloxacillin"].endswith("— Cloxacillin")
    assert titles["3.2.P.1"] == "Description and Composition of Drug Product"


def test_slug_is_stable_and_path_safe():
    assert slugify_subject("Ampicillin Trihydrate") == "ampicillin-trihydrate"
    assert slugify_subject("Amoxicillin/Clavulanate") == "amoxicillin-clavulanate"


# ---- placement --------------------------------------------------------------


def test_each_substance_gets_its_own_ctd_folder():
    a = folder_for_section_instance("3.2.S.4.1", "ampicillin")
    c = folder_for_section_instance("3.2.S.4.1", "cloxacillin")
    assert a != c
    assert a == (
        "m3/32-body-data/32s/32s-ampicillin/32s4-control-of-drug-substance/32s41-specification"
    )


def test_unmapped_repeating_section_fails_loudly():
    """Same instinct as folder_for_section: a section with no mapped folder
    must fail at build time, not land in a guessed location no reviewer
    would look in."""
    with pytest.raises(KeyError, match="3.2.S.9"):
        folder_for_section_instance("3.2.S.9", "ampicillin")


# ---- content ----------------------------------------------------------------


def _render(project, number, subject):
    storage = InMemoryStorageClient()
    result = render_section(number, project, storage=storage, subject=subject)
    return Document(io.BytesIO(storage.get(result.storage_key)))


def test_specification_renders_its_own_substances_rows():
    project = _load(build_ampiclox)
    cloxacillin = {a.inn_name: a for a in project.product.apis}["Cloxacillin"]

    doc = _render(project, "3.2.S.4.1", cloxacillin)

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Cloxacillin" in text
    assert "Ampicillin" not in text  # this copy is about ONE substance
    rows = doc.tables[0].rows
    # header + one row per test on file
    assert len(rows) == 1 + len(cloxacillin.specification)
    body = "\n".join(c.text for r in rows for c in r.cells)
    assert "Assay (anhydrous basis)" in body
    assert "90.0" in body or "98.0" in body


def test_general_information_shows_only_its_own_substances_structure():
    """The 2.3 QOS shows every active's structure; 3.2.S.1 must show exactly
    one, or a combination product's drug-substance section would repeat the
    whole product in each copy."""
    project = _load(build_ampiclox)
    ampicillin = {a.inn_name: a for a in project.product.apis}["Ampicillin"]

    doc = _render(project, "3.2.S.1", ampicillin)

    assert len(doc.inline_shapes) == 1
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Structural formula — Ampicillin:" in text
    assert "Cloxacillin" not in text


def test_nomenclature_omits_identifiers_the_substance_does_not_have():
    """A table of blank rows reads as missing data; omitting the row reads
    as not applicable, which is what an absent CEP actually means."""
    project = _load(build_ampiclox)
    ampicillin = {a.inn_name: a for a in project.product.apis}["Ampicillin"]

    doc = _render(project, "3.2.S.1", ampicillin)

    body = "\n".join(c.text for r in doc.tables[0].rows for c in r.cells)
    assert "Salt / hydrate form as manufactured" in body
    assert "CEP number" not in body  # not on file for this fixture


# ---- the DTD's own requirements ---------------------------------------------


def test_r17_flags_an_active_with_no_manufacturer():
    """`m3-2-s-drug-substance` declares `manufacturer` #REQUIRED, so this is
    the DTD's rule, not ours. Caught at the data layer while the applicant
    can still fix it, rather than at a gateway."""
    project = _load(build_ampiclox)
    project.product.apis[0].manufacturer = None

    findings = [f for f in run_all(project).findings if f.rule_id == "R17"]

    assert len(findings) == 1
    assert "Ampicillin" in findings[0].message


def test_backbone_info_refuses_to_invent_a_manufacturer():
    """Rather than substituting a placeholder: a backbone naming the wrong
    maker of an API is worse than a failed build."""
    project = _load(build_ampiclox)
    project.product.apis[1].manufacturer = None

    with pytest.raises(ValueError, match="Cloxacillin"):
        drug_substance_info(project)


def test_r07_reads_specification_rows_per_substance():
    project = _load(build_ampiclox)
    by_name = {a.inn_name: a for a in project.product.apis}
    by_name["Cloxacillin"].specification = []

    findings = [f for f in run_all(project).findings if f.rule_id == "R07"]

    assert len(findings) == 1
    assert "Cloxacillin" in findings[0].message


def test_specification_rows_render_in_declared_order():
    """Specification tables are conventionally ordered and an assessor reads
    release against stability side by side, so the order is content."""
    project = _load(build_examox)
    api = project.product.apis[0]
    api.specification = [
        SpecificationTest(test_name="Second", method="m", acceptance_criterion="c", sort_order=1),
        SpecificationTest(test_name="First", method="m", acceptance_criterion="c", sort_order=0),
    ]

    doc = _render(project, "3.2.S.4.1", api)

    names = [r.cells[0].text for r in doc.tables[0].rows[1:]]
    assert names == ["First", "Second"]
