"""Tests: the real docxtpl template engine (P04) renders section .docx files
from the EXAMOX seed data, with data slots correctly filled and narrative
slots left as clearly-marked placeholders when no narrative is supplied."""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.storage import InMemoryStorageClient
from app.models import Base
from app.seed.ampiclox import build_ampiclox
from app.seed.examox import build_examox
from app.templating.registry import get_section
from app.templating.render import render_section


@pytest.fixture
def project():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)
    project = build_examox(buggy=False)
    session.add(project)
    session.commit()
    session.refresh(project)
    return project


@pytest.fixture
def fdc_project():
    """AMPICLOX -- the fixed-dose combination fixture. Anything in the
    template engine that quietly assumes one active per product fails here
    and nowhere else, since every other seed happens to be single-API."""
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)
    project = build_ampiclox(buggy=False)
    session.add(project)
    session.commit()
    session.refresh(project)
    return project


def _document_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_get_section_rejects_unknown_number():
    with pytest.raises(KeyError):
        get_section("9.9.9")


def test_render_cover_letter_fills_data_slots(project):
    storage = InMemoryStorageClient()
    result = render_section("1.0", project, storage=storage)

    assert result.storage_key == f"projects/{project.id}/sections/1.0.docx"
    assert result.size_bytes == len(storage.get(result.storage_key))

    text = _document_text(storage.get(result.storage_key))
    assert "NAFDAC" in text
    assert "EXAMOX" in text
    assert "Exagon" in text
    assert "[[AI DRAFT PENDING" in text  # narrative slot, no narrative supplied
    assert "{{" not in text and "{%" not in text  # no leftover jinja tags


def test_render_p1_fills_composition_table(project):
    storage = InMemoryStorageClient()
    result = render_section("3.2.P.1", project, storage=storage)

    text = _document_text(storage.get(result.storage_key))
    assert "EXAMOX" in text
    assert "Amoxicillin 500 mg" in text  # per-API strength_display
    assert "144.00" in text  # declared batch quantity, from BatchFormulaLine
    assert "[[AI DRAFT PENDING" in text
    assert "{{" not in text and "{%" not in text


def test_render_stability_summary_with_narrative_supplied(project):
    storage = InMemoryStorageClient()
    conclusion = "Supports the declared 24-month shelf life per ICH Q1A(R2)."
    result = render_section(
        "3.2.P.8.1", project, narrative={"conclusion": conclusion}, storage=storage
    )

    text = _document_text(storage.get(result.storage_key))
    assert "24 months" in text  # declared shelf life
    assert "30C/65%RH" in text  # stability study condition, from StabilityStudy row
    assert conclusion in text
    assert "[[AI DRAFT PENDING" not in text  # narrative was supplied, no placeholder


def test_rendered_docx_is_a_valid_zip(project):
    storage = InMemoryStorageClient()
    result = render_section("3.2.P.1", project, storage=storage)
    # a .docx is a zip archive -- this is the cheapest possible check that
    # docxtpl produced a well-formed document, not just non-empty bytes.
    with zipfile.ZipFile(io.BytesIO(storage.get(result.storage_key))) as zf:
        assert "word/document.xml" in zf.namelist()


def test_render_qos_embeds_the_structure_image_when_smiles_present(project):
    storage = InMemoryStorageClient()
    result = render_section(
        "2.3", project, narrative={"overview": "A beta-lactam antibiotic."}, storage=storage
    )

    doc = Document(io.BytesIO(storage.get(result.storage_key)))
    assert len(doc.inline_shapes) == 1  # the EXAMOX seed's API has a smiles set
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "A beta-lactam antibiotic." in text
    assert "Structure not available" not in text
    # every structure is captioned with its substance, single-API included --
    # an uncaptioned formula is one an assessor can't tie to a named active.
    assert "Structural formula — Amoxicillin:" in text


def test_render_qos_embeds_one_structure_per_active_for_a_combination(fdc_project):
    """The regression this file exists to prevent: the QOS structure slot
    used to render `apis[0]` only, so AMPICLOX's summary showed ampicillin
    and silently dropped cloxacillin. 2.3.S is repeated per drug substance,
    so an FDC owes one captioned structural formula per active."""
    storage = InMemoryStorageClient()

    result = render_section("2.3", fdc_project, storage=storage)

    doc = Document(io.BytesIO(storage.get(result.storage_key)))
    assert len(doc.inline_shapes) == 2
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Structural formula — Ampicillin:" in text
    assert "Structural formula — Cloxacillin:" in text


def test_render_qos_placeholder_is_per_active_not_all_or_nothing(fdc_project):
    """A half-entered combination degrades per substance: the active that
    has a SMILES still gets its picture, and only the one that doesn't gets
    the placeholder -- rather than one missing field blanking both."""
    by_name = {api.inn_name: api for api in fdc_project.product.apis}
    by_name["Cloxacillin"].smiles = None
    storage = InMemoryStorageClient()

    result = render_section("2.3", fdc_project, storage=storage)

    doc = Document(io.BytesIO(storage.get(result.storage_key)))
    assert len(doc.inline_shapes) == 1
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Structural formula — Ampicillin:" in text
    assert "Structural formula — Cloxacillin:" in text
    assert text.count("Structure not available") == 1


def test_render_qos_shows_placeholder_when_smiles_missing(project):
    project.product.apis[0].smiles = None
    storage = InMemoryStorageClient()

    result = render_section("2.3", project, storage=storage)

    doc = Document(io.BytesIO(storage.get(result.storage_key)))
    assert len(doc.inline_shapes) == 0
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Structure not available" in text
    assert "[[AI DRAFT PENDING" in text
