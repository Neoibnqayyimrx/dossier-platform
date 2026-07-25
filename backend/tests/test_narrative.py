"""Tests for the P05 narrative generation service: the full retrieve ->
prompt -> generate -> guardrail -> audit pipeline, human review gating, and
its wiring into the P04 template context.

Needs the real `pg_session_factory` (see tests/conftest.py) because
`generate_narrative` calls P03's `search()`, which compiles to pgvector's
cosine_distance operator -- no SQLite equivalent, same reason
test_knowledge.py's search tests need it. Skips itself if no Postgres is
reachable.
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.core.storage import InMemoryStorageClient
from app.knowledge.embeddings import FakeEmbeddingClient
from app.knowledge.ingest import ingest_document
from app.llm.client import FakeLLMClient, LLMClient
from app.models.enums import NarrativeStatus
from app.narrative.context import get_approved_narrative
from app.narrative.generate import generate_narrative
from app.narrative.guardrails import NarrativeGuardrailError
from app.narrative.review import approve_narrative, edit_narrative
from app.seed.examox import build_examox
from app.templating.render import render_section

EMBEDDER = FakeEmbeddingClient()

Q1A_EXCERPT = """\
2.1.6. Stability Testing
For long term studies, frequency of testing should be sufficient to
establish the stability profile of the drug substance under the declared
storage condition.
"""


class _FabricatingLLMClient(LLMClient):
    """Cites a source that was never retrieved -- for exercising the
    citation guardrail's hard-block path."""

    def generate(self, system: str, messages: list[dict[str, str]], max_tokens: int = 1024) -> str:
        return "Per compendial requirements. [Source: USP 43 2020]"


class _LeakyLLMClient(LLMClient):
    """Invents a figure absent from the facts -- for exercising the
    numeric-leakage warning path."""

    def generate(self, system: str, messages: list[dict[str, str]], max_tokens: int = 1024) -> str:
        return "This product remains stable for 999 months under all conditions."


async def _seed_project_and_kb(db):
    project = build_examox(buggy=False)
    db.add(project)
    await db.commit()

    await ingest_document(
        db,
        EMBEDDER,
        source="ICH",
        title="Q1A(R2) Stability Testing",
        version="Step 4, 2003-02-06",
        license="ich-harmonised-guideline",
        text=Q1A_EXCERPT,
        chunk_max_chars=500,
    )
    return project


async def test_generate_returns_grounded_prose_with_sources_and_audit_row(pg_session_factory):
    async with pg_session_factory() as db:
        project = await _seed_project_and_kb(db)

        result = await generate_narrative(
            db,
            project,
            "3.2.P.8.1",
            "conclusion",
            embedding_client=EMBEDDER,
            llm_client=FakeLLMClient(),
        )

        assert result.narrative.status == NarrativeStatus.PENDING
        assert result.narrative.final_text is None
        assert "EXAMOX" in result.narrative.output  # a data fact, from the prompt
        assert "Stability Testing" in result.narrative.output  # the retrieved source excerpt
        assert len(result.narrative.sources) >= 1
        # NOTE: not asserting zero leakage warnings here -- the fake LLM
        # echoes the whole prompt, including the retrieved source excerpt's
        # own section numbers ("2.1.6", etc.), which aren't part of the
        # product facts and so legitimately trip the leakage heuristic. The
        # guardrail's own behavior (warn vs. block, and on which numbers)
        # is covered directly by test_guardrails.py and the two tests below.


async def test_generate_rejects_unknown_slot(pg_session_factory):
    async with pg_session_factory() as db:
        project = await _seed_project_and_kb(db)
        with pytest.raises(ValueError, match="not a narrative slot"):
            await generate_narrative(
                db,
                project,
                "3.2.P.8.1",
                "not_a_real_slot",
                embedding_client=EMBEDDER,
                llm_client=FakeLLMClient(),
            )


async def test_fabricated_citation_is_blocked_and_nothing_is_persisted(pg_session_factory):
    async with pg_session_factory() as db:
        project = await _seed_project_and_kb(db)

        with pytest.raises(NarrativeGuardrailError, match="USP 43 2020"):
            await generate_narrative(
                db,
                project,
                "3.2.P.8.1",
                "conclusion",
                embedding_client=EMBEDDER,
                llm_client=_FabricatingLLMClient(),
            )

        approved = await get_approved_narrative(db, project.id, "3.2.P.8.1")
        assert approved == {}


async def test_numeric_leakage_is_flagged_but_not_blocking(pg_session_factory):
    async with pg_session_factory() as db:
        project = await _seed_project_and_kb(db)

        result = await generate_narrative(
            db,
            project,
            "3.2.P.8.1",
            "conclusion",
            embedding_client=EMBEDDER,
            llm_client=_LeakyLLMClient(),
        )

        assert result.narrative.status == NarrativeStatus.PENDING  # still persisted
        assert any("999" in w for w in result.warnings)
        assert result.narrative.warnings and "999" in result.narrative.warnings


async def test_only_approved_narrative_reaches_the_rendered_section(pg_session_factory):
    async with pg_session_factory() as db:
        project = await _seed_project_and_kb(db)
        result = await generate_narrative(
            db,
            project,
            "3.2.P.8.1",
            "conclusion",
            embedding_client=EMBEDDER,
            llm_client=FakeLLMClient(),
        )
        narrative = result.narrative

        # PENDING: nothing usable yet, and the rendered doc shows the placeholder.
        approved = await get_approved_narrative(db, project.id, "3.2.P.8.1")
        assert approved == {}

        storage = InMemoryStorageClient()
        pending_render = render_section("3.2.P.8.1", project, narrative=approved, storage=storage)
        pending_text = _document_text(storage.get(pending_render.storage_key))
        assert "[[AI DRAFT PENDING" in pending_text

        approve_narrative(narrative)
        await db.commit()

        approved = await get_approved_narrative(db, project.id, "3.2.P.8.1")
        assert approved == {"conclusion": narrative.output}

        approved_render = render_section("3.2.P.8.1", project, narrative=approved, storage=storage)
        approved_text = _document_text(storage.get(approved_render.storage_key))
        assert narrative.output in approved_text
        assert "[[AI DRAFT PENDING" not in approved_text


async def test_edited_narrative_overrides_the_original_output(pg_session_factory):
    async with pg_session_factory() as db:
        project = await _seed_project_and_kb(db)
        result = await generate_narrative(
            db,
            project,
            "3.2.P.8.1",
            "conclusion",
            embedding_client=EMBEDDER,
            llm_client=FakeLLMClient(),
        )
        narrative = result.narrative

        edit_narrative(narrative, "Hand-edited conclusion text, reviewed by QA.")
        await db.commit()

        approved = await get_approved_narrative(db, project.id, "3.2.P.8.1")
        assert approved == {"conclusion": "Hand-edited conclusion text, reviewed by QA."}
        assert narrative.status == NarrativeStatus.EDITED


def _document_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
