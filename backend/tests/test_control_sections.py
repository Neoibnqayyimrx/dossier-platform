"""Tests: one specification model, three owners; batches; impurities (P20).

The phase's claim is that a specification is ONE artifact the CTD asks for
of three different things, and that building it three times is how the
three come to disagree. These tests pin that claim from both ends:

  * the same model, the same renderer and the same rules serve a drug
    substance, an excipient and the finished product;
  * a batch result outside its own specification's limit blocks the export
    and says which batch, which test, which result and which limit.

And one negative: P13's drug-substance behaviour must be exactly what it
was. The migration widened the table; it did not move anything.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.storage import InMemoryStorageClient
from app.ctd.structure import folder_for_section, folder_for_section_instance
from app.models import (
    Base,
    BatchAnalysis,
    BatchAnalysisResult,
    Impurity,
    ImpurityType,
    SpecificationOwnerKind,
    SpecificationTest,
)
import app.validation.rules  # noqa: F401  registers rules
from app.seed.ampiclox import build_ampiclox
from app.seed.examox import build_examox
from app.templating.instances import expand_sections, repeat_element_info
from app.templating.render import render_section
from app.validation.acceptance import evaluate
from app.validation.engine import Severity, run_all


def _load(builder, **kwargs):
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)
    project = builder(**kwargs)
    session.add(project)
    session.commit()
    session.refresh(project)
    return project


def _render(project, number, subject=None):
    storage = InMemoryStorageClient()
    result = render_section(number, project, storage=storage, subject=subject)
    return Document(io.BytesIO(storage.get(result.storage_key)))


def _table_text(doc, index=0):
    return "\n".join(cell.text for row in doc.tables[index].rows for cell in row.cells)


# ---- the owner is polymorphic, and the database enforces it -----------------


def test_one_specification_model_answers_all_three_owners():
    project = _load(build_examox, buggy=False)
    product = project.product

    kinds = {
        product.apis[0].specification[0].owner_kind,
        product.excipients[0].specification[0].owner_kind,
        product.specification[0].owner_kind,
    }

    assert kinds == {
        SpecificationOwnerKind.DRUG_SUBSTANCE,
        SpecificationOwnerKind.EXCIPIENT,
        SpecificationOwnerKind.DRUG_PRODUCT,
    }
    # One TABLE, not three: the rows all live in specification_test.
    assert {type(row) for row in product.specification} == {SpecificationTest}


def test_the_database_refuses_a_specification_row_with_two_owners():
    """The CHECK constraint, not the application, is what makes this true.

    A row with two owners would render into two sections with the same id
    and raise nothing anywhere; a row with none would silently vanish from
    every section. Only the database can refuse both regardless of which
    code path wrote the row -- including a migration or a psql session.
    """
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)
    project = build_examox(buggy=False)
    session.add(project)
    session.commit()

    session.add(
        SpecificationTest(
            active_ingredient_id=project.product.apis[0].id,
            product_id=project.product.id,
            test_name="Assay",
            method="HPLC",
            acceptance_criterion="98.0 - 102.0 %",
        )
    )
    with pytest.raises(IntegrityError):
        session.commit()


def test_the_database_refuses_a_specification_row_with_no_owner():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)

    session.add(
        SpecificationTest(test_name="Assay", method="HPLC", acceptance_criterion="98.0 - 102.0 %")
    )
    with pytest.raises(IntegrityError):
        session.commit()


# ---- ... and renders correctly into each of the three sections --------------


@pytest.mark.parametrize(
    "number, subject_from, expected_owner_label",
    [
        ("3.2.S.4.1", lambda p: p.product.apis[0], "Drug substance"),
        ("3.2.P.4.1", lambda p: p.product.excipients[0], "Excipient"),
        ("3.2.P.5.1", lambda p: None, "Drug product"),
    ],
)
def test_the_same_specification_renders_into_each_section(
    number, subject_from, expected_owner_label
):
    """One template, three sections. The heading and the owner label come
    from the context, which is the mechanism that lets the three share
    `specification.docx` at all."""
    project = _load(build_examox, buggy=False)
    subject = subject_from(project)
    owner = subject if subject is not None else project.product

    doc = _render(project, number, subject)

    text = "\n".join(p.text for p in doc.paragraphs)
    assert number in text
    assert expected_owner_label in text
    rows = doc.tables[0].rows
    # header + one row per test on file, for whichever owner this is
    assert len(rows) == 1 + len(owner.specification)


def test_the_drug_product_specification_is_not_the_drug_substances():
    """The concrete reason the owner had to become polymorphic rather than
    the drug substance's table being reused: a purified API is held to
    98.0-102.0 %, a finished dosage form to 90.0-110.0 % of label claim,
    and the finished product has tests (dissolution, uniformity of dosage
    units) that a drug substance does not have at all."""
    project = _load(build_examox, buggy=False)

    substance = _table_text(_render(project, "3.2.S.4.1", project.product.apis[0]))
    product = _table_text(_render(project, "3.2.P.5.1"))

    assert "98.0 - 102.0 % w/w" in substance
    assert "90.0 - 110.0 % of label claim" in product
    assert "Dissolution" in product
    assert "Dissolution" not in substance


def test_each_excipient_gets_its_own_3_2_p_4_1():
    """The excipient repeat axis, declared in P19 and unusable until the
    specification could belong to an excipient."""
    project = _load(build_examox, buggy=False)

    keys = [i.key for i in expand_sections(project)]

    assert "3.2.P.4.1-starch" in keys
    assert "3.2.P.4.1-magnesium-stearate" in keys
    # ...each in its own folder, so an assessor can tell which material a
    # folder holds without opening it.
    assert folder_for_section_instance("3.2.P.4.1", "starch").endswith(
        "32p4-control-of-excipients/excipient-starch/32p41-specification"
    )
    # 3.2.P.4.2 covers every excipient in ONE document, so it does not
    # repeat and keeps a plain folder.
    assert "3.2.P.4.2" in keys
    assert folder_for_section("3.2.P.4.2").endswith("32p42-analytical-procedures")


def test_an_excipient_section_names_only_its_own_excipient():
    project = _load(build_examox, buggy=False)
    starch = {e.name: e for e in project.product.excipients}["Starch"]

    doc = _render(project, "3.2.P.4.1", starch)

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Starch" in text
    assert "Magnesium Stearate" not in text


# ---- the rule the whole design exists for -----------------------------------


def test_an_out_of_specification_batch_result_blocks_the_export():
    """R22, and the four values a finding must name.

    The limit is not a copy of the specification's limit -- it is reached
    through `result.specification_test`, which is why tightening 3.2.S.4.1
    re-judges every batch already on file.
    """
    project = _load(build_ampiclox, buggy=True)

    report = run_all(project)
    errors = [f for f in report.errors() if f.rule_id == "R22"]

    assert errors, "the buggy AMPICLOX fixture plants an out-of-specification assay"
    message = errors[0].message
    assert "API/AMP/24/0021" in message  # the batch
    assert "Assay (anhydrous basis)" in message  # the test
    assert "103.4 % w/w" in message  # the result
    assert "98.0 - 102.0 % w/w" in message  # the limit
    assert not report.is_exportable()


def test_a_compliant_dossier_raises_no_out_of_specification_error():
    project = _load(build_ampiclox, buggy=False)

    errors = [f for f in run_all(project).errors() if f.rule_id == "R22"]

    assert errors == []


def test_tightening_the_specification_re_judges_batches_already_on_file():
    """The single-source-of-truth claim, tested rather than asserted: no
    result is re-entered, and every batch is re-checked."""
    project = _load(build_ampiclox, buggy=False)
    api = project.product.apis[0]
    assay = {row.test_name: row for row in api.specification}["Assay (anhydrous basis)"]

    assert [f for f in run_all(project).errors() if f.rule_id == "R22"] == []
    assay.acceptance_criterion = "99.5 - 100.0 % w/w"

    errors = [f for f in run_all(project).errors() if f.rule_id == "R22"]

    assert errors
    assert "99.5 - 100.0 % w/w" in errors[0].message


def test_an_uncheckable_result_is_reported_as_unchecked_not_as_passed():
    """The distinction the parser exists to preserve. An acceptance
    criterion nothing can parse must never be silently credited as met."""
    project = _load(build_examox, buggy=False)
    api = project.product.apis[0]
    batch = api.batch_analyses[0]
    test = batch.results[0].specification_test
    test.acceptance_criterion = "Conforms to the reference chromatogram"
    batch.results[0].result = "See attached chromatogram"

    findings = [f for f in run_all(project).findings if f.rule_id == "R22"]
    unchecked = [f for f in findings if f.severity is Severity.INFO]

    assert unchecked, "an unparseable pair must be reported, not passed over"
    assert test.test_name in unchecked[0].message
    # ...and it is not an ERROR: the platform does not know it failed.
    assert not [f for f in findings if f.severity is Severity.ERROR]


@pytest.mark.parametrize(
    "criterion, result, expected",
    [
        ("98.0 - 102.0 % w/w", "99.4 %", True),
        ("98.0 - 102.0 % w/w", "103.4 %", False),
        ("98.0 - 102.0 % w/w", "97.9 %", False),
        ("NMT 1.0 %", "0.42 %", True),
        ("NMT 1.0 %", "1.4 %", False),
        ("NLT 80 % (Q) in 45 minutes", "94 % in 45 min", True),
        ("NLT 80 % (Q) in 45 minutes", "74 % in 45 min", False),
        ("Complies", "Complies", True),
        # The ordering bug this parser is shaped to avoid: "does not
        # comply" CONTAINS "comply", so a naive substring test in the wrong
        # order reads a failure as a pass.
        ("Complies", "Does not comply", False),
        # A stated failure outranks the numbers: a result reading "Fails"
        # against a range must not be rescued by having no parseable number.
        ("98.0 - 102.0 %", "Fails", False),
        # Neither parseable: reported as undecidable, never as passed.
        ("White to off-white powder", "Slightly yellow powder", None),
        # A European decimal comma must not read as a hundredfold error.
        ("NMT 0,15 %", "0,12 %", True),
        ("NMT 0,15 %", "0,20 %", False),
    ],
)
def test_acceptance_criteria_are_read_the_way_a_pharmacist_reads_them(criterion, result, expected):
    assert evaluate(criterion, result) is expected


# ---- results cannot answer a test that is not in the specification ----------


def test_a_result_is_bound_to_the_test_it_answers():
    """Not a test NAME typed again -- a foreign key. Renaming the test
    cannot orphan the result, and the limit can never be a stale copy."""
    project = _load(build_examox, buggy=False)
    api = project.product.apis[0]
    result = api.batch_analyses[0].results[0]

    result.specification_test.test_name = "Assay (renamed)"

    assert result.specification_test.test_name == "Assay (renamed)"
    findings = [f for f in run_all(project).findings if f.rule_id == "R22"]
    assert all("Not tested" not in f.message for f in findings)


# ---- impurities -------------------------------------------------------------


def test_impurities_render_per_substance_and_for_the_product():
    project = _load(build_ampiclox, buggy=False)
    ampicillin = {a.inn_name: a for a in project.product.apis}["Ampicillin"]

    substance = _table_text(_render(project, "3.2.S.3.2", ampicillin))
    product = _table_text(_render(project, "3.2.P.5.5"))

    assert "Impurity A (6-aminopenicillanic acid)" in substance
    assert "process-related" in substance
    # 3.2.P.5.5 is degradation-only, which is the regulatory point of the
    # section: formulating cannot introduce a process impurity of the API.
    assert "process-related" not in product
    assert "Total degradation products" in product


def test_r11_reaches_impurity_limits_and_skips_ich_thresholds():
    """The reminder can only exist because `limit_source` is data. It must
    fire on a monograph-derived limit and NOT on an ICH threshold -- the
    latter does not change with a pharmacopoeial edition, and a reminder to
    go check one would be noise that teaches the reader to ignore R11."""
    project = _load(build_examox, buggy=False)

    messages = [f.message for f in run_all(project).findings if f.rule_id == "R11"]
    impurity_reminders = [m for m in messages if m.startswith("Impurity ")]

    assert any("Impurity A (6-aminopenicillanic acid)" in m for m in impurity_reminders)
    assert any("BP monograph" in m for m in impurity_reminders)
    assert not any("ICH Q3A" in m for m in impurity_reminders)
    assert not any("ICH Q3C" in m for m in impurity_reminders)


# ---- the eCTD backbone: a second repeating ELEMENT ---------------------------


def test_excipients_repeat_as_a_heading_element_like_drug_substances():
    """`m3-2-p-4-control-of-excipients` is starred in the ICH DTD with an
    `excipient` attribute -- structurally the same shape as
    `m3-2-s-drug-substance`, reached from the other end of Module 3. Two
    excipients must therefore be two sibling ELEMENTS, not two leaves
    merged under one heading (the P13 `_Node` bug, arriving by a new
    route)."""
    project = _load(build_examox, buggy=False)

    info = repeat_element_info(project)

    excipient_axes = {key: attrs for key, (axis, attrs) in info.items() if axis == "excipient"}
    assert "3.2.P.4.1-starch" in excipient_axes
    assert excipient_axes["3.2.P.4.1-starch"] == {"excipient": "Starch"}
    # The drug substance axis is unchanged and still carries both of the
    # DTD's #REQUIRED attributes.
    substance_axes = {key: attrs for key, (axis, attrs) in info.items() if axis == "drug_substance"}
    assert set(substance_axes["3.2.S.4.1-amoxicillin"]) == {"substance", "manufacturer"}


# ---- P13 did not regress ----------------------------------------------------


def test_the_migration_widened_the_table_without_moving_a_row():
    """P20's fourth definition-of-done, stated as a test rather than as a
    promise in a migration docstring: every drug-substance specification
    row still hangs off its active ingredient, exactly as P13 left it."""
    project = _load(build_ampiclox, buggy=False)

    for api in project.product.apis:
        assert api.specification, "every active still has its own specification"
        for row in api.specification:
            assert row.active_ingredient_id == api.id
            assert row.product_id is None
            assert row.excipient_id is None
            assert row.owner_kind is SpecificationOwnerKind.DRUG_SUBSTANCE


def test_batch_and_impurity_rows_reject_an_excipient_owner():
    """`BatchAnalysis` and `Impurity` declare the NARROWER two-owner space
    on purpose: the CTD has no excipient batch-analysis or impurity leaf,
    and a row the dossier could never render should not be storable."""
    assert SpecificationOwnerKind.EXCIPIENT not in BatchAnalysis.__owner_kinds__
    assert SpecificationOwnerKind.EXCIPIENT not in Impurity.__owner_kinds__
    assert not hasattr(BatchAnalysis, "excipient_id")
    assert not hasattr(Impurity, "excipient_id")


def test_a_batch_analysis_section_puts_the_limit_beside_the_result():
    """The layout claim in batch_analysis_context: a limit printed in one
    table and the numbers judged against it in another is what lets an
    out-of-specification result pass unnoticed."""
    project = _load(build_examox, buggy=False)
    api = project.product.apis[0]

    doc = _render(project, "3.2.S.4.4", api)

    # tables[0] is the batch list; tables[1] is the results.
    results = doc.tables[1]
    header = [cell.text for cell in results.rows[0].cells]
    assert header == ["Test", "Acceptance criterion", "Batch", "Result"]
    body = _table_text(doc, index=1)
    assert "API/AMO/24/0011" in body
    assert "98.0 - 102.0 % w/w" in body


def test_a_batch_result_carries_its_specification_order():
    """So the batch table and the specification table are read ACROSS each
    other rather than against each other."""
    project = _load(build_examox, buggy=False)
    batch = project.product.apis[0].batch_analyses[0]

    orders = [result.sort_order for result in batch.results]

    assert orders == sorted(orders)
    for result in batch.results:
        assert result.sort_order == result.specification_test.sort_order


def test_a_batch_with_no_results_says_so_rather_than_rendering_blank():
    project = _load(build_examox, buggy=False)
    api = project.product.apis[0]
    api.batch_analyses = []

    doc = _render(project, "3.2.S.4.4", api)

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "NO BATCH ANALYSIS ON FILE" in text


def test_analytical_procedures_separates_compendial_from_in_house():
    """Where the specification is pharmacopoeial the procedure is a
    citation; where it is in-house it needs a description. Both occur in
    one table, which is why the section is hybrid."""
    project = _load(build_examox, buggy=False)

    doc = _render(project, "3.2.P.5.2")

    body = _table_text(doc)
    assert "Compendial" in body
    # The seeded drug-product specification includes one in-house method.
    assert "In-house" in body
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "In-house method AM-014" in text


def test_a_fully_compendial_specification_reproduces_no_method_text():
    """AGENTS.md §5 at the template layer: with no in-house method there is
    nothing to describe, and the section says so rather than inviting prose
    that would have to come from a copyrighted monograph."""
    project = _load(build_examox, buggy=False)

    doc = _render(project, "3.2.S.4.2", project.product.apis[0])

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "No in-house method is used" in text


def test_a_justification_prints_the_limits_it_is_defending():
    project = _load(build_examox, buggy=False)

    doc = _render(project, "3.2.P.5.6")

    body = _table_text(doc)
    assert "90.0 - 110.0 % of label claim" in body


def test_batch_results_and_impurities_survive_a_round_trip():
    """A plain persistence check: these are the first two models with a
    two-level owned collection, and the cascade has to reach the leaves."""
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)
    project = build_examox(buggy=False)
    session.add(project)
    session.commit()

    api = project.product.apis[0]
    reloaded = session.get(type(api), api.id)

    assert reloaded.batch_analyses
    assert all(batch.results for batch in reloaded.batch_analyses)
    assert reloaded.impurities
    assert isinstance(reloaded.impurities[0], Impurity)
    assert isinstance(reloaded.impurities[0].impurity_type, ImpurityType)
    assert isinstance(reloaded.batch_analyses[0].results[0], BatchAnalysisResult)
