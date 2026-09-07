"""Tests: bioequivalence as data, not a paragraph (P22).

The phase's claim is that leaf 5.3.1.2 carries the entire scientific
argument for a multisource approval, and that while a bioequivalence study
was `kind` + `reference_product` + a summary, the only checkable fact about
it was that a row existed. These tests pin the claim from both ends:

  * a 90 % confidence interval outside the acceptance window BLOCKS the
    export and NAMES the parameter and the bound;
  * a study run against a comparator the application does not declare is
    caught -- the cross-module check the platform was built for;
  * 1.4.1 renders from study data with no hand-entered content, and
    CANNOT disagree with 5.2, because both read the same rows.

Plus the machinery that makes the biowaiver decision real: answering the
conditional question is what puts the request in the package, and rule R06
refuses two routes or none.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.storage import InMemoryStorageClient
from app.ctd.region_profiles import NAFDAC_PROFILE
from app.ctd.structure import folder_for_section
from app.models import (
    Base,
    BEDoseRegimen,
    BEFedState,
    BEStudyDesign,
    BioequivalenceResult,
    BioequivalenceStudy,
    Biowaiver,
    BiowaiverKind,
    PKParameter,
    ReferenceProduct,
)
import app.validation.rules  # noqa: F401  registers rules
from app.seed.ampiclox import build_ampiclox
from app.seed.examox import build_examox
from app.templating.instances import expand_sections
from app.templating.registry import SECTIONS
from app.templating.render import render_section
from app.validation.engine import Severity, run_all


def _load(builder, **kwargs):
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)
    project = builder(**kwargs)
    session.add(project)
    session.commit()
    session.refresh(project)
    return project


def _render(project, number):
    storage = InMemoryStorageClient()
    result = render_section(number, project, storage=storage)
    return Document(io.BytesIO(storage.get(result.storage_key)))


def _text(doc) -> str:
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    cells = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return f"{paragraphs}\n{cells}"


def _in_vivo(project) -> None:
    """Answer the conditional biowaiver leaves as an in vivo filing.

    Written out rather than left unanswered: R06 needs a route, and an
    unanswered condition is silence -- which is the state R19 reports,
    not a decision.
    """
    project.condition_answers = {**project.condition_answers, "1.2.17": False, "1.2.18": False}


# ---- the study is data now --------------------------------------------------


def test_a_study_carries_its_design_its_comparator_and_its_intervals():
    """What `ClinicalEntry` could not hold. Every one of these is a fact an
    assessor reads off the BTI form, and none of them survives being a
    sentence."""
    project = _load(build_examox, buggy=False)
    study = project.product.bioequivalence_studies[0]

    assert study.design is BEStudyDesign.CROSSOVER
    assert study.fed_state is BEFedState.FASTING
    assert study.dose_regimen is BEDoseRegimen.SINGLE_DOSE
    assert study.dropouts == 2
    assert {result.parameter for result in study.results} == set(PKParameter)

    # The comparator is a REFERENCE, not a string: it has a batch and an
    # expiry, which is what makes "was it in date when it was dosed?"
    # answerable at all.
    assert study.reference_product is not None
    assert study.reference_product.batch_number
    assert study.reference_product.expiry_date > study.completion_date

    # The test batch is the batch 3.2.P.5.4 files, by foreign key.
    assert study.test_batch is not None
    assert study.test_batch in project.product.batch_analyses


def test_the_results_are_ordered_the_way_every_study_report_prints_them():
    """Cmax, then AUC(0-t), then AUC(0-inf) -- rate before extent. Sorting
    by the stored string would put AUC first and make every generated
    table disagree with every study report."""
    study = _load(build_examox, buggy=False).product.bioequivalence_studies[0]
    assert [r.parameter for r in study.results] == [
        PKParameter.CMAX,
        PKParameter.AUC_0_T,
        PKParameter.AUC_0_INF,
    ]


# ---- R25: the interval, and the window that judges it -----------------------


def test_r25_blocks_the_export_and_names_the_parameter_and_the_bound():
    """The definition of done, and the rule the phase exists for.

    AMPICLOX's buggy variant plants a Cmax lower bound of 76.40 % against
    an 80.00 % floor -- the realistic failure, since a generic that
    dissolves slightly slower than the innovator fails on rate long before
    it fails on extent.
    """
    project = _load(build_ampiclox)  # buggy=True by default
    report = run_all(project)

    findings = [f for f in report.findings if f.rule_id == "R25"]
    assert findings, "an out-of-window confidence interval must be reported"
    assert all(f.severity is Severity.ERROR for f in findings)
    assert not report.is_exportable()

    message = findings[0].message
    # Names the PARAMETER and the BOUND. "The study failed" sends a filer
    # back to a hundred-page report; this is a line they can act on.
    assert "Cmax" in message
    assert "lower" in message
    assert "76.40" in message
    assert "80.00 - 125.00 %" in message


def test_r25_is_silent_on_a_study_that_passes():
    project = _load(build_examox, buggy=False)
    assert not [f for f in run_all(project).findings if f.rule_id == "R25"]


def test_the_acceptance_window_is_config_and_narrows_for_a_narrow_index_drug():
    """The window is a REGULATORY PARAMETER, so it lives in the region
    profile -- and which of the two applies is a property of the molecule,
    so the product carries the flag. Neither is a branch in the rule."""
    project = _load(build_examox, buggy=False)
    product = project.product

    assert NAFDAC_PROFILE.window_for(product).label == "80.00 - 125.00 %"

    # An interval of 85.00-105.00 is bioequivalent for an ordinary drug and
    # is NOT for a narrow-therapeutic-index one. Same study, same numbers,
    # different answer -- which is the whole point of putting the window in
    # config rather than in the rule.
    cmax = next(
        r for r in product.bioequivalence_studies[0].results if r.parameter is PKParameter.CMAX
    )
    cmax.ci_lower, cmax.ci_upper = Decimal("85.00"), Decimal("105.00")
    assert not [f for f in run_all(project).findings if f.rule_id == "R25"]

    product.narrow_therapeutic_index = True
    assert NAFDAC_PROFILE.window_for(product).label == "90.00 - 111.11 %"

    findings = [f for f in run_all(project).findings if f.rule_id == "R25"]
    assert any("Cmax" in f.message and "lower" in f.message for f in findings)
    assert any("90.00 - 111.11 %" in f.message for f in findings)


# ---- R26: the cross-module comparator check ---------------------------------


def test_r26_catches_a_study_run_against_a_comparator_the_application_does_not_declare():
    """The definition of done's second half, and the cross-module check the
    platform was built for.

    AMPICLOX's buggy variant declares "Ampiclox 250/250 mg capsules" on the
    application and dosed "Ospamox 500 mg capsules" in the study -- the
    entirely ordinary failure where the comparator originally planned is
    not the one the CRO could source.
    """
    project = _load(build_ampiclox)
    findings = [f for f in run_all(project).findings if f.rule_id == "R26"]

    assert findings
    assert findings[0].severity is Severity.ERROR
    assert "Ospamox" in findings[0].message
    assert "Ampiclox" in findings[0].message
    assert not run_all(project).is_exportable()


def test_r26_forgives_spacing_and_case_but_not_a_different_brand():
    """A check that fires on every filing is a check the filer learns to
    ignore. "Amoxil 500mg Capsules" and "Amoxil 500 mg capsules" are one
    product; no normalisation makes Amoxil into Ospamox."""
    project = _load(build_examox, buggy=False)
    product = project.product
    study = product.bioequivalence_studies[0]

    study.reference_product.name = "  AMOXIL   500mg   CAPSULES "
    assert not [f for f in run_all(project).findings if f.rule_id == "R26"]

    study.reference_product.name = "Ospamox 500 mg capsules"
    assert [f for f in run_all(project).findings if f.rule_id == "R26"]


# ---- R27: the cross-module batch-size check ---------------------------------


def test_r27_flags_a_biobatch_too_small_to_represent_production():
    """A genuine regulatory finding that exists only in the space BETWEEN
    two modules: the test batch size is in Module 5, the commercial batch
    size is in 3.2.P.3.2. Nobody reading either alone can see it.

    WHO TRS 992 Annex 7: a tenth of the commercial batch, or 100 000
    units, whichever is greater. The seeds declare a 250 000-unit
    commercial batch, so the 100 000 floor binds.
    """
    project = _load(build_examox, buggy=False)
    study = project.product.bioequivalence_studies[0]

    assert not [f for f in run_all(project).findings if f.rule_id == "R27"]

    study.test_batch_size_units = 8_000
    findings = [f for f in run_all(project).findings if f.rule_id == "R27"]
    assert findings
    assert findings[0].severity is Severity.ERROR
    assert "8,000" in findings[0].message
    assert "250,000" in findings[0].message
    assert "100,000" in findings[0].message


def test_r27_is_silent_when_the_batch_size_has_not_been_entered():
    """A data-entry gap is not a scale defect, and a rule that guessed here
    would block filings over an empty field."""
    project = _load(build_examox, buggy=False)
    project.product.bioequivalence_studies[0].test_batch_size_units = None
    assert not [f for f in run_all(project).findings if f.rule_id == "R27"]


# ---- R06: exactly one route -------------------------------------------------


def test_r06_blocks_a_filing_with_no_bioequivalence_route_at_all():
    project = _load(build_examox, buggy=False)
    project.product.bioequivalence_studies.clear()
    findings = [f for f in run_all(project).findings if f.rule_id == "R06"]
    assert any("no bioequivalence route" in f.message for f in findings)


def test_r06_blocks_a_biowaiver_claimed_alongside_an_in_vivo_study():
    """They are ALTERNATIVES. Both is a contradiction: the application says
    at once that a human study was necessary and that it was not."""
    project = _load(build_examox, buggy=False)
    project.condition_answers = {"1.2.17": True}
    project.product.biowaivers.append(
        Biowaiver(kind=BiowaiverKind.BCS_BASED, strength="500 mg", bcs_class=1)
    )

    findings = [f for f in run_all(project).findings if f.rule_id == "R06"]
    assert any("mutually exclusive" in f.message for f in findings)
    assert not run_all(project).is_exportable()


def test_r06_accepts_a_biowaiver_on_its_own():
    """A filing with no human study is not an incomplete dossier -- it is
    the other route, and the platform must not demand both."""
    project = _load(build_examox, buggy=False)
    project.product.bioequivalence_studies.clear()
    project.condition_answers = {"1.2.17": True, "1.2.18": False}
    project.product.biowaivers.append(
        Biowaiver(kind=BiowaiverKind.BCS_BASED, strength="500 mg", bcs_class=1)
    )
    assert not [f for f in run_all(project).findings if f.rule_id == "R06"]


def test_r06_refuses_a_claim_with_no_request_data_behind_it():
    """The leaf IS in the package -- the answer is what put it there -- so
    it would ship as a request naming no strength and citing nothing."""
    project = _load(build_examox, buggy=False)
    project.product.bioequivalence_studies.clear()
    project.condition_answers = {"1.2.17": True, "1.2.18": False}
    findings = [f for f in run_all(project).findings if f.rule_id == "R06"]
    assert any("no biowaiver request data" in f.message for f in findings)


def test_r06_does_not_ask_a_new_chemical_entity_for_a_comparator():
    """An NCE has nothing to be equivalent to. Before P22 this rule fired
    on every project regardless, which was wrong in a way nothing noticed
    because nothing had built an NCE filing yet."""
    from app.models.enums import SubmissionType

    project = _load(build_examox, buggy=False)
    project.product.bioequivalence_studies.clear()
    project.submission_type = SubmissionType.NEW_CHEMICAL_ENTITY
    assert not [f for f in run_all(project).findings if f.rule_id == "R06"]


# ---- 1.4.1: a Module 1 form built from Module 5 numbers ---------------------


def test_the_bti_form_has_no_narrative_slot_at_all():
    """The target TOC's own note on this leaf: "Derived entirely from the
    BE study data. No prose." A slot here could state an interval the
    results table disproves."""
    assert SECTIONS["1.4.1"].narrative_slots == []


def test_the_bti_form_prints_the_study_data_and_the_window_it_is_judged_against():
    project = _load(build_examox, buggy=False)
    _in_vivo(project)
    text = _text(_render(project, "1.4.1"))

    study = project.product.bioequivalence_studies[0]
    assert study.study_identifier in text
    assert study.cro_name in text
    assert study.reference_product.batch_number in text
    assert "80.00 - 125.00 %" in text
    # The conclusion is STATED, not left to the reader: a form printing a
    # failing interval without saying it fails has to be re-checked by hand.
    assert "bioequivalence is demonstrated" in text.lower()


def test_the_bti_form_and_the_tabular_listing_cannot_disagree():
    """The definition of done's third claim, and the whole premise of the
    platform in one assertion.

    1.4.1 is a Module 1 form and 5.2 is a Module 5 table. In a
    hand-assembled dossier they are two separate typings of one fact,
    done weeks apart, and an assessor reads them against each other.
    Here neither has any independent existence: change the study and both
    change together, because both read the same rows.
    """
    project = _load(build_examox, buggy=False)
    _in_vivo(project)

    study = project.product.bioequivalence_studies[0]
    bti = _text(_render(project, "1.4.1"))
    listing = _text(_render(project, "5.2"))

    # The BTI form breaks the comparator out into named fields (an agency
    # reads a form field by field); the listing prints it as one column.
    # What has to agree is the FACTS, and they do because both read the
    # same row.
    for text in (bti, listing):
        assert study.study_identifier in text
        assert study.reference_product.name in text
        assert study.reference_product.manufacturer in text

    # Now break the study, and watch BOTH documents follow. A dossier in
    # which only one of them moved is the defect this phase removes.
    study.results[0].ci_lower = Decimal("70.00")
    bti_after = _text(_render(project, "1.4.1"))
    listing_after = _text(_render(project, "5.2"))

    assert "70.00" in bti_after
    assert "NOT demonstrated" in bti_after
    assert "NOT demonstrated" in listing_after


def test_the_tabular_listing_is_walked_from_what_is_actually_filed():
    """5.2's note in the target TOC is the whole specification: "Must not
    be able to disagree with what is actually filed under 5.3.1."
    """
    project = _load(build_examox, buggy=False)
    listing = _text(_render(project, "5.2"))
    assert "5.3.1.2" in listing  # the leaf the report is filed at

    project.product.bioequivalence_studies.clear()
    after = _text(_render(project, "5.2"))
    assert "EXAMOX/BE/2025-01" not in after


# ---- placement and the leaves themselves ------------------------------------


def test_every_new_section_has_somewhere_to_be_filed():
    """`folder_for_section` raises rather than guessing, and a Module 1
    leaf is placed by the region profile instead. A registered section
    with no home is a document that cannot be built."""
    assert folder_for_section("5.2").startswith("m5/")
    assert folder_for_section("5.3.1.2").startswith("m5/")

    module1 = {slot.section_number: slot.folder for slot in NAFDAC_PROFILE.module1_slots}
    for number in ("1.4.1", "1.2.17", "1.2.18"):
        assert number in module1, f"{number} has no Module 1 slot"


def test_the_study_summary_ships_beside_the_uploaded_report_not_instead_of_it():
    """5.3.1.2 holds TWO leaves: the CRO's study report and the structured
    summary generated from the data.

    Assembly lets an uploaded file win over a rendered one at the same key
    -- correctly, since a generated stand-in for a signed document is not
    an improvement on it. `leaf_suffix` is what keeps the summary out of
    that collision; without it, attaching the report would silently delete
    the summary, which is exactly the "looks complete, is not" failure the
    platform exists to prevent.
    """
    project = _load(build_examox, buggy=False)
    _in_vivo(project)
    instances = {i.key: i for i in expand_sections(project)}

    assert "5.3.1.2-summary" in instances
    assert "5.3.1.2" not in instances
    assert instances["5.3.1.2-summary"].number == "5.3.1.2"


def test_a_biowaiver_leaf_is_in_the_package_only_when_it_is_claimed():
    """P17's machinery earning its keep. Answering "yes" to 1.2.17 is not
    a preference recorded somewhere -- it is a document appearing in
    Module 1."""
    project = _load(build_examox, buggy=False)

    project.condition_answers = {"1.2.17": False, "1.2.18": False}
    assert "1.2.17" not in {i.key for i in expand_sections(project)}

    project.condition_answers = {"1.2.17": True, "1.2.18": False}
    assert "1.2.17" in {i.key for i in expand_sections(project)}


def test_the_biowaiver_request_renders_its_evidence_and_names_its_leaf():
    project = _load(build_examox, buggy=False)
    project.condition_answers = {"1.2.17": True, "1.2.18": False}
    project.product.biowaivers.append(
        Biowaiver(
            kind=BiowaiverKind.BCS_BASED,
            strength="500 mg",
            bcs_class=1,
            dissolution_similarity_f2=Decimal("62.40"),
        )
    )
    text = _text(_render(project, "1.2.17"))

    assert "500 mg" in text
    assert "Class 1" in text
    # The verdict is printed beside the number, not just the number.
    assert "62.40" in text and "similar" in text


def test_a_biowaiver_knows_which_leaf_it_is_filed_at():
    """One place, so the renderer, the rule and the route control cannot
    each decide for themselves."""
    assert Biowaiver(kind=BiowaiverKind.BCS_BASED, strength="5 mg").section_number == "1.2.17"
    assert (
        Biowaiver(kind=BiowaiverKind.ADDITIONAL_STRENGTH, strength="10 mg").section_number
        == "1.2.18"
    )


# ---- the scale finding ------------------------------------------------------


def test_a_biowaiver_records_its_strength_as_text_because_a_filing_is_one_strength():
    """The phase's scale finding, asserted rather than only written down.

    A Project points at ONE Product, and a Product's strength lives on its
    ActiveIngredient rows. There is no product family, so a filing cannot
    span 5 mg and 10 mg -- and leaf 1.2.18, "biowaiver request for an
    ADDITIONAL strength", presumes exactly that span. The platform can
    record the claim and print the request; it cannot cross-check the
    other strength's composition, because the other strength is not in the
    filing.

    This test exists so that the day `Product` grows a family, it fails --
    and whoever is holding it is pointed at the leaf that was waiting.
    """
    project = _load(build_examox, buggy=False)
    product = project.product

    # One strength per active, and no way to say "this product also comes
    # in 250 mg".
    assert not hasattr(product, "strengths")
    assert not hasattr(product, "family")
    assert {api.strength_value for api in product.apis} == {Decimal("500.000")}

    waiver = Biowaiver(kind=BiowaiverKind.ADDITIONAL_STRENGTH, strength="250 mg")
    assert isinstance(waiver.strength, str)


# ---- the API ----------------------------------------------------------------


async def test_results_are_replaced_as_a_set_not_a_row_at_a_time(auth_client):
    """A study report states the three intervals together, off one page. A
    study holding two of its three is worse than one holding none, because
    it looks answered."""
    product = (
        await auth_client.post(
            "/products", json={"brand_name": "TESTOX", "generic_name": "Testolol"}
        )
    ).json()

    created = await auth_client.post(
        f"/products/{product['id']}/bioequivalence",
        json={
            "study_identifier": "TESTOX/BE/01",
            "design": "crossover",
            "fed_state": "fasting",
            "dose_regimen": "single dose",
        },
    )
    assert created.status_code == 201
    study_id = created.json()["id"]

    rows = [
        {"parameter": p, "ci_lower": "95.00", "ci_upper": "105.00"}
        for p in ("Cmax", "AUC(0-t)", "AUC(0-inf)")
    ]
    saved = await auth_client.put(f"/bioequivalence/{study_id}/results", json=rows)
    assert saved.status_code == 200
    assert [row["parameter"] for row in saved.json()] == ["Cmax", "AUC(0-t)", "AUC(0-inf)"]

    # Replace, not merge: a parameter the filer removed has to be gone
    # from 1.4.1 and 5.2 as well.
    replaced = await auth_client.put(f"/bioequivalence/{study_id}/results", json=rows[:1])
    assert [row["parameter"] for row in replaced.json()] == ["Cmax"]


async def test_a_transposed_confidence_interval_is_refused(auth_client):
    """A bound the wrong way round is almost always a transposed paste, and
    it matters: swapped, an interval that really fails could be reported as
    passing."""
    product = (
        await auth_client.post(
            "/products", json={"brand_name": "TESTOX", "generic_name": "Testolol"}
        )
    ).json()
    study_id = (
        await auth_client.post(
            f"/products/{product['id']}/bioequivalence",
            json={
                "study_identifier": "TESTOX/BE/02",
                "design": "crossover",
                "fed_state": "fasting",
                "dose_regimen": "single dose",
            },
        )
    ).json()["id"]

    response = await auth_client.put(
        f"/bioequivalence/{study_id}/results",
        json=[{"parameter": "Cmax", "ci_lower": "125.00", "ci_upper": "80.00"}],
    )
    assert response.status_code == 422
    assert "lower bound" in response.text


def test_a_study_and_its_comparator_can_be_built_without_a_database():
    """The transient object graph the rules actually read. Kept as a test
    because every rule in this file is exercised through one."""
    reference = ReferenceProduct(name="Amoxil 500 mg capsules", manufacturer="GSK")
    study = BioequivalenceStudy(
        study_identifier="X/1",
        design=BEStudyDesign.PARALLEL,
        fed_state=BEFedState.FED,
        dose_regimen=BEDoseRegimen.MULTIPLE_DOSE,
        reference_product=reference,
        start_date=date(2025, 1, 1),
    )
    study.results.append(
        BioequivalenceResult(
            parameter=PKParameter.CMAX,
            ci_lower=Decimal("79.99"),
            ci_upper=Decimal("101.00"),
        )
    )
    assert reference.identity == "Amoxil 500 mg capsules (GSK)"
    assert study.design_summary == "parallel, multiple dose, fed"
    assert study.results[0].outside(Decimal("80.00"), Decimal("125.00")) == ["lower"]
