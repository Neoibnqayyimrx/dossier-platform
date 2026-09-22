"""A validation report a person can read (gap Phase 5b).

Until this phase findings existed only as JSON (`FindingRead`), which is
fine for the web UI and useless for the conversation that actually happens
around a dossier: a regulatory lead forwarding "here is where we stand" to
a QA reviewer, or filing the state of a package beside the package. This
renders the SAME findings the readiness endpoint and the eCTD validation
endpoint return -- it adds no judgement of its own -- as a PDF.

WHY a PDF through the platform's own pipeline (python-docx, then
`convert_docx_to_pdf`) rather than HTML or a new PDF library: it is the
artifact you can attach to an email, it needed no new dependency, and it is
exactly how the table of contents is already produced (app/ctd/toc.py),
including the two table courtesies that page learnt the hard way -- the
header row repeats on every page, and a row never splits across a break.

WHY waived checks get their own section, above the findings: a package
built over a waived ERROR is byte-for-byte as convincing as one that
passed. The only place that difference can surface is in a report like
this one, next to the reason a human recorded for it.
"""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.shared import Inches

from app.assembly.pdf import convert_docx_to_pdf
from app.ctd.toc import keep_row_on_one_page, repeat_header_on_every_page
from app.validation.engine import Finding, Severity

# Most consequential first: a reader looks at the top of the list.
_SEVERITY_ORDER = (Severity.ERROR, Severity.WARNING, Severity.INFO, Severity.ADVISORY)

_SOURCE_LABELS = {
    "data-rule": "data rule",
    "mechanical-ectd": "eCTD check",
    "external-validator": "external validator",
    "ai-reviewer": "AI reviewer",
}


def build_validation_report_pdf(
    *,
    project,
    scope: str,
    findings: list[Finding],
    is_exportable: bool,
    waived: list[tuple[str, str]],
    generated_at: datetime,
) -> bytes:
    """The report as PDF bytes.

    `scope` says which layers the findings came from ("Data rules" or "eCTD
    sequence 0001 -- all validation layers"). `waived` is (rule id, the
    reason a human recorded). `generated_at` is passed in rather than read
    from the clock so a caller -- and a test -- controls it: a report is a
    statement about a moment, and it says which one.
    """
    waived_ids = {rule_id for rule_id, _ in waived}
    doc = Document()
    doc.add_heading(f"Validation report -- {project.product.brand_name}", level=1)

    for label, value in (
        ("Project", project.name),
        ("Region", project.region.value),
        ("Scope", scope),
        ("Generated", generated_at.strftime("%Y-%m-%d %H:%M UTC")),
        ("Result", "Exportable" if is_exportable else "Blocked -- unresolved ERROR findings"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)

    counts = [
        f"{sum(1 for f in findings if f.severity is severity)} {severity.value}"
        for severity in _SEVERITY_ORDER
    ]
    doc.add_paragraph("Findings: " + ", ".join(counts))

    if waived:
        doc.add_heading("Waived checks", level=2)
        doc.add_paragraph(
            "A waived check is a recorded human decision, not a pass: the package "
            "was allowed to build despite it. The reason below is the one given "
            "when the waiver was recorded."
        )
        _table(doc, ("Rule", "Reason recorded"), sorted(waived), widths=(0.9, 5.4))

    if not findings:
        doc.add_paragraph("No findings.")
    for severity in _SEVERITY_ORDER:
        of_severity = [f for f in findings if f.severity is severity]
        if not of_severity:
            continue
        doc.add_heading(f"{severity.value} ({len(of_severity)})", level=2)
        _table(
            doc,
            ("Rule", "Section", "Category", "Source", "Finding"),
            widths=(0.9, 0.8, 1.3, 0.95, 2.35),
            rows=[
                (
                    # On its own line: squeezed onto the id's line it broke
                    # before the closing bracket.
                    f.rule_id + ("\n(waived)" if f.rule_id in waived_ids else ""),
                    f.section or "--",
                    f.category,
                    _SOURCE_LABELS.get(f.source, f.source),
                    f.message,
                )
                for f in sorted(of_severity, key=lambda f: (f.rule_id, f.section or ""))
            ],
        )

    doc.add_paragraph(
        "ERROR blocks export unless a human waives it with a recorded reason. "
        "WARNING and INFO never block. ADVISORY findings -- the AI reviewer's, "
        "and the note that no agency validator ran -- never block by construction."
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    return convert_docx_to_pdf(buffer.getvalue(), bookmark_title="Validation report")


def _table(doc, headers: tuple[str, ...], rows, *, widths: tuple[float, ...]) -> None:
    """A table with fixed column `widths` (inches), header repeated per page.

    WHY fixed widths: left to share the page equally, the Finding column --
    the one that matters -- wrapped after every word while the rule id sat
    in a column as wide as itself. Word keeps widths on the grid AND on each
    cell, and LibreOffice honours the cells, so both are set.
    """
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    repeat_header_on_every_page(table.rows[0])
    for values in rows:
        row = table.add_row()
        keep_row_on_one_page(row)
        for cell, value in zip(row.cells, values):
            cell.text = str(value)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
        for cell in column.cells:
            cell.width = Inches(width)
