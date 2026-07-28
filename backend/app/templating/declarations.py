"""Generated-content document generator for administrative declarations
(P08). See app.models.declaration for why this is NOT the same shape as
app.templating.certificates's placeholder: a Declaration's text is fully
knowable from data on file (Applicant + Product + Project) -- what's
missing is a human's wet signature and, for some types, a notary's seal.
So the generated document is real, reviewable content with a clearly
marked action-required block, not a "we don't have this file" stand-in.

WHY plain python-docx, not docxtpl: like certificates.py, there's no
pre-authored template with `{{ }}` slots -- the whole document (boilerplate
text + a handful of data values) is built directly in Python.

WHY the storage key is scoped by project, not product: a Declaration
(e.g. Power of Attorney) names a representative for THIS filing -- the
same product filed again later could have a different one. See
app.models.declaration's docstring.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from docx import Document

from app.core.storage import StorageClient, get_storage_client
from app.models.declaration import Declaration
from app.models.enums import DECLARATIONS_REQUIRING_NOTARIZATION, DeclarationType
from app.models.project import Project

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_MISSING = "[[NOT YET ON FILE]]"

# Boilerplate body text per declaration type -- the part of a NAFDAC
# administrative declaration that's the same for every applicant, only
# ever needing the applicant/product facts filled in around it.
_BODY_TEXT: dict[DeclarationType, str] = {
    DeclarationType.POWER_OF_ATTORNEY: (
        "The undersigned, being duly authorized to act on behalf of the "
        "Applicant, hereby appoints the named representative below to act "
        "as the Applicant's agent for all matters relating to this "
        "registration filing with NAFDAC, including submission, "
        "correspondence, and receipt of decisions."
    ),
    DeclarationType.DECLARATION_OF_AUTHENTICITY: (
        "The undersigned hereby declares that all information, data, and "
        "documents submitted in support of this application are true, "
        "accurate, and complete to the best of the Applicant's knowledge."
    ),
    DeclarationType.GMP_COMPLIANCE_UNDERTAKING: (
        "The undersigned undertakes that the manufacturing site(s) named in "
        "this dossier will maintain Good Manufacturing Practice compliance "
        "throughout the validity of this registration, and will notify "
        "NAFDAC of any material change in manufacturing arrangements."
    ),
}


@dataclass
class DeclarationRenderResult:
    declaration_id: str
    storage_key: str
    size_bytes: int


def render_declaration(
    declaration: Declaration, project: Project, storage: StorageClient | None = None
) -> DeclarationRenderResult:
    applicant = project.applicant
    product = project.product

    doc = Document()
    doc.add_heading(declaration.declaration_type.value.replace("-", " ").title(), level=1)
    doc.add_paragraph(f"Applicant: {applicant.company_name if applicant else _MISSING}")
    doc.add_paragraph(
        f"Product: {product.brand_name} ({product.generic_name}) {product.strength_display}"
    )
    doc.add_paragraph(_BODY_TEXT[declaration.declaration_type])
    if applicant and applicant.authorized_representative_name:
        doc.add_paragraph(
            f"Representative: {applicant.authorized_representative_name}, "
            f"{applicant.authorized_representative_title or ''}".rstrip(", ")
        )
    needs_notarization = declaration.declaration_type in DECLARATIONS_REQUIRING_NOTARIZATION
    doc.add_paragraph(
        "ACTION REQUIRED: this document must be printed, signed by an "
        "authorized signatory"
        + (", and notarized/legalized" if needs_notarization else "")
        + " before this dossier is submitted. An unsigned copy must never be filed."
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()

    storage = storage or get_storage_client()
    key = f"projects/{project.id}/declarations/{declaration.id}.docx"
    storage.put(key, data, DOCX_CONTENT_TYPE)

    return DeclarationRenderResult(
        declaration_id=str(declaration.id), storage_key=key, size_bytes=len(data)
    )
