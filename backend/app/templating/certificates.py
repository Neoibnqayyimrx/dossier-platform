"""Placeholder document generator for regulatory certificates (P04).

WHY a placeholder, never generated content: a `Certificate` (CPP, GMP, CEP,
CoA, free-sale -- see app.models.certificate) is proof issued by a third
party: a regulator, EDQM, a testing lab. The platform cannot write one --
pretending otherwise would be exactly the fabrication AGENTS.md §5's
determinism boundary forbids. What it CAN do is guarantee the assembled
package has a clearly-labeled slot at the right place, so a missing
certificate is a loud, visible gap a human sees immediately -- not a
silent one only discovered during compilation review.

WHY plain python-docx, not docxtpl: there is no template with `{{ }}`
slots to fill here, just a handful of Python values arranged on a page --
docxtpl's whole purpose (Jinja substitution into pre-authored template
XML) doesn't apply, so reaching for it would just be indirection.

WHY the storage key is scoped by product, not project: a Certificate hangs
off `Product` (master data), not `Project` (a specific filing) --
app.models.project's own docstring notes the same product can be filed as
several Projects over time. A GMP certificate for a manufacturing site
doesn't change per filing, so its placeholder lives alongside the product,
not nested under whichever project happens to render it first.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

from docx import Document

from app.core.storage import StorageClient, get_storage_client
from app.models.certificate import Certificate

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass
class CertificateRenderResult:
    certificate_id: str
    storage_key: str
    size_bytes: int


def render_certificate_placeholder(
    certificate: Certificate, storage: StorageClient | None = None
) -> CertificateRenderResult:
    doc = Document()
    doc.add_heading("PLACEHOLDER — REPLACE THIS FILE", level=1)
    doc.add_paragraph(
        f"This is a placeholder for a {certificate.certificate_type.value} certificate."
    )
    doc.add_paragraph(f"Issuing authority: {certificate.issuing_authority or '(not yet known)'}")
    doc.add_paragraph(
        f"Certificate number: {certificate.certificate_number or '(not yet obtained)'}"
    )
    doc.add_paragraph(f"Issue date: {certificate.issue_date or '(not yet obtained)'}")
    doc.add_paragraph(f"Expiry date: {certificate.expiry_date or '(not yet obtained)'}")
    doc.add_paragraph(
        "ACTION REQUIRED: replace this file with the actual certificate document "
        "before this dossier is submitted. This placeholder must never be filed "
        "as-is."
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()

    storage = storage or get_storage_client()
    key = f"products/{certificate.product_id}/certificates/{certificate.id}.docx"
    storage.put(key, data, DOCX_CONTENT_TYPE)

    return CertificateRenderResult(
        certificate_id=str(certificate.id), storage_key=key, size_bytes=len(data)
    )
